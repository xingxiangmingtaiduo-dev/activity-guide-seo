from pathlib import Path
import re, shutil, subprocess
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from PIL import Image

ROOT=Path(__file__).parent
OUT=ROOT/'outputs'/'2026-10-12'; OUT.mkdir(parents=True,exist_ok=True)
SOURCE='https://allthings.how/anime-vanguards-2nd-anniversary-update-features-and-release-status/'
IMGURL='https://static.allthings.how/wp-content/uploads/2026/09/anime-vanguards-2nd-anniversary-update-is-a-new-game-1025012901.webp'
# Date-qualified status, preparation steps and product uses in every target language.
DATA={
'en':[
'According to the reference updated on September 16, 2026, the anniversary update was still in production, with no official release date. Community countdowns were speculation. Check the latest official announcement before planning your session.',
'Memoria customization is planned. Save Memoria Shards until the final system is explained.',
'The reference describes these options in version 14.5. Check that they remain available in your game.',
['Try Lunar Bridge: 11 floors, each with a demon boss. Bosses can drop Demon Familiars for Demon Controller units.', 'Use the in-game Codex to look up units. Gauntlets offers an 8-player bracket; Odyssey offers card choices between runs.', 'Save Trait Rerolls, Memoria Shards and upgrade materials. Wait for the damage rebalance and final Familiar passives before spending heavily.'],
['Run compatible Android games in the cloud to reduce local storage, battery and hardware demands during longer sessions.', 'Use remote access to return to your session and check progress. You still need to play; cloud hosting does not itself automate farming.', 'Test login and gameplay first. Your connection, server node and game compatibility affect the experience.'],
'Save first, choose upgrades after the final details arrive. Try a short UGphone session before planning a longer run.',
'Transcendent Sung Jin-woo preview cutscene', 'Image source', 'Preparation guide', 'Save resources and review the announced anniversary features before choosing upgrades.'],
'id':[
'Menurut referensi yang diperbarui pada 16 September 2026, update ulang tahun masih dibuat dan belum memiliki tanggal rilis resmi. Hitung mundur komunitas hanya perkiraan. Cek pengumuman resmi terbaru sebelum mengatur jadwal bermain.',
'Kustomisasi Memoria direncanakan. Simpan Memoria Shards sampai sistem akhirnya dijelaskan.',
'Referensi menyebut pilihan berikut dalam versi 14.5. Cek apakah masih tersedia di dalam game.',
['Coba Lunar Bridge: 11 lantai, masing-masing dengan bos iblis. Bos dapat menjatuhkan Demon Familiar untuk unit Demon Controller.', 'Gunakan Codex dalam game untuk mencari informasi unit. Gauntlets memakai bagan turnamen 8 pemain; Odyssey memberi pilihan kartu di antara putaran.', 'Simpan Trait Rerolls, Memoria Shards, dan bahan peningkatan. Tunggu penyesuaian damage dan efek pasif Familiar final sebelum memakai banyak bahan.'],
['Jalankan game Android yang kompatibel di cloud untuk mengurangi beban penyimpanan, baterai, dan perangkat lokal saat bermain lebih lama.', 'Akses jarak jauh membantu kamu kembali ke sesi dan memeriksa progres. Kamu tetap perlu bermain; hosting cloud tidak otomatis melakukan farming.', 'Uji login dan gameplay terlebih dahulu. Koneksi, node server, dan kompatibilitas game memengaruhi pengalaman.'],
'Kumpulkan bahan dulu, pilih peningkatan setelah detail final keluar. Coba sesi singkat di UGphone sebelum merencanakan sesi panjang.',
'Cuplikan adegan Transcendent Sung Jin-woo', 'Sumber gambar', 'Panduan persiapan', 'Simpan bahan dan pelajari fitur ulang tahun yang diumumkan sebelum memilih peningkatan.'],
'pt':[
'Segundo a referência atualizada em 16 de setembro de 2026, a atualização de aniversário ainda estava em produção e não tinha data oficial. As contagens regressivas da comunidade eram especulação. Consulte o anúncio oficial mais recente antes de planejar uma sessão.',
'Está prevista a personalização de Memoria. Guarde Memoria Shards até a explicação do sistema final.',
'A referência descreve estas opções na versão 14.5. Verifique se continuam disponíveis no jogo.',
['Experimente Lunar Bridge: 11 andares, cada um com um chefe demônio. Os chefes podem deixar Demon Familiars para unidades Demon Controller.', 'Consulte unidades no Codex do jogo. Gauntlets tem um torneio de 8 jogadores; Odyssey oferece escolhas de cartas entre partidas.', 'Guarde Trait Rerolls, Memoria Shards e materiais de melhoria. Espere o rebalanceamento de dano e os efeitos passivos finais dos Familiars antes de gastar muito.'],
['Execute jogos Android compatíveis na nuvem para reduzir o uso de armazenamento, bateria e recursos do aparelho local em sessões longas.', 'O acesso remoto permite voltar à sessão e verificar o progresso. Você ainda precisa jogar; a hospedagem na nuvem não automatiza a coleta.', 'Teste o login e a jogabilidade primeiro. A conexão, o servidor utilizado e a compatibilidade do jogo afetam a experiência.'],
'Junte materiais primeiro e escolha melhorias após os detalhes finais. Teste uma sessão curta no UGphone antes de planejar uma sessão longa.',
'Cena da prévia de Transcendent Sung Jin-woo', 'Fonte da imagem', 'Guia de preparação', 'Guarde materiais e confira as novidades anunciadas para o aniversário antes de escolher melhorias.'],
'es':[
'Según la referencia actualizada el 16 de septiembre de 2026, la actualización del aniversario seguía en desarrollo y no tenía fecha oficial. Las cuentas atrás de la comunidad eran especulaciones. Consulta el anuncio oficial más reciente antes de planificar una sesión.',
'Está prevista la personalización de Memoria. Guarda Memoria Shards hasta conocer el sistema final.',
'La referencia describe estas opciones en la versión 14.5. Comprueba si siguen disponibles en el juego.',
['Prueba Lunar Bridge: 11 plantas, cada una con un jefe demonio. Los jefes pueden soltar Demon Familiars para unidades Demon Controller.', 'Consulta las unidades en el Codex del juego. Gauntlets ofrece un torneo de 8 jugadores; Odyssey permite elegir cartas entre partidas.', 'Guarda Trait Rerolls, Memoria Shards y materiales de mejora. Espera al reajuste de daño y a los efectos pasivos finales de los Familiars antes de gastar mucho.'],
['Ejecuta juegos Android compatibles en la nube para reducir el uso de almacenamiento, batería y recursos del dispositivo local durante sesiones largas.', 'El acceso remoto permite volver a tu sesión y revisar el progreso. Sigues teniendo que jugar; alojar el juego en la nube no automatiza el farmeo.', 'Prueba primero el acceso y la jugabilidad. La conexión, el nodo del servidor y la compatibilidad del juego afectan a la experiencia.'],
'Reúne materiales primero y elige mejoras cuando lleguen los detalles finales. Prueba una sesión corta en UGphone antes de planificar una larga.',
'Escena del adelanto de Transcendent Sung Jin-woo', 'Fuente de la imagen', 'Guía de preparación', 'Guarda materiales y revisa las novedades anunciadas del aniversario antes de elegir mejoras.'],
'th':[
'ตามบทความอ้างอิงที่อัปเดตวันที่ 16 กันยายน 2026 อัปเดตครบรอบยังอยู่ระหว่างพัฒนาและไม่มีวันเปิดตัวอย่างเป็นทางการ เคาน์ดาวน์ของชุมชนเป็นเพียงการคาดเดา ควรตรวจสอบประกาศทางการล่าสุดก่อนวางแผนเล่น',
'มีแผนเพิ่มการปรับแต่ง Memoria ควรเก็บ Memoria Shards ไว้จนกว่าจะทราบรายละเอียดระบบจริง',
'บทความอ้างอิงกล่าวถึงตัวเลือกเหล่านี้ในเวอร์ชัน 14.5 ตรวจสอบในเกมว่ายังเปิดให้เล่นหรือไม่',
['ลอง Lunar Bridge ซึ่งมี 11 ชั้นและบอสปีศาจในแต่ละชั้น บอสมีโอกาสดรอป Demon Familiar สำหรับยูนิต Demon Controller', 'ค้นหาข้อมูลยูนิตใน Codex ภายในเกม Gauntlets เป็นการแข่งขันแบบสายแข่ง 8 คน ส่วน Odyssey ให้เลือกการ์ดระหว่างรอบ', 'เก็บ Trait Rerolls, Memoria Shards และวัสดุอัปเกรด รอผลปรับสมดุลความเสียหายและค่าพาสซีฟ Familiar จริงก่อนใช้ทรัพยากรจำนวนมาก'],
['รันเกม Android ที่รองรับบนคลาวด์ เพื่อลดภาระพื้นที่เก็บข้อมูล แบตเตอรี่ และเครื่องหลักเมื่อเล่นนาน', 'ใช้การเข้าถึงระยะไกลเพื่อกลับเข้าเซสชันและตรวจสอบความคืบหน้า คุณยังต้องเล่นเอง การโฮสต์บนคลาวด์ไม่ได้ฟาร์มให้อัตโนมัติ', 'ทดสอบการล็อกอินและการเล่นก่อน เครือข่าย โหนดเซิร์ฟเวอร์ และความเข้ากันได้ของเกมมีผลต่อประสบการณ์'],
'เก็บวัสดุก่อนแล้วค่อยเลือกอัปเกรดเมื่อรายละเอียดครบ ลองเล่นบน UGphone ช่วงสั้น ๆ ก่อนวางแผนเล่นนาน',
'ฉากจากตัวอย่าง Transcendent Sung Jin-woo', 'แหล่งที่มาของภาพ', 'คู่มือเตรียมตัว', 'เก็บวัสดุและตรวจสอบฟีเจอร์ครบรอบที่ประกาศก่อนเลือกอัปเกรด'],
'zh-TW':[
'依據 2026 年 9 月 16 日更新的參考文章，二週年改版當時仍在製作中，尚無官方上線日期。社群倒數僅屬猜測。安排遊玩前，請先確認最新官方公告。',
'預計加入 Memoria 自訂功能，先保留 Memoria Shards，等正式系統說明再決定用途。',
'參考文章介紹了 14.5 版本的以下玩法，請先在遊戲中確認目前是否仍可參加。',
['挑戰 Lunar Bridge：共 11 層，每層都有惡魔首領。首領可掉落供 Demon Controller 單位使用的 Demon Familiar。', '用遊戲內 Codex 查詢單位資料。Gauntlets 採 8 人淘汰賽，Odyssey 則在每局之間提供卡牌選擇。', '保存 Trait Rerolls（特性重抽）、Memoria Shards 與升級材料，等傷害平衡和 Familiar 被動效果定案後，再集中投入。'],
['在雲端執行相容的 Android 遊戲，長時間遊玩時可減少本機效能、儲存空間與電量負擔。', '遠端存取方便回到遊戲並查看進度。仍須自行操作，雲端託管本身不會自動刷本。', '先測試登入與遊玩。實際體驗受網路、伺服器節點及遊戲相容性影響。'],
'先存材料，再依正式改版選擇升級方向。想安排長時間遊玩，可先在 UGphone 試玩一小段時間。',
'Transcendent Sung Jin-woo 預告過場畫面', '圖片來源', '改版準備攻略', '整理二週年預告內容與資源準備建議，等正式平衡公布後再決定升級方向。'],
'vi':[
'Theo bài tham khảo cập nhật ngày 16 tháng 9 năm 2026, bản cập nhật kỷ niệm vẫn đang được phát triển và chưa có ngày ra mắt chính thức. Đếm ngược của cộng đồng chỉ là dự đoán. Hãy xem thông báo chính thức mới nhất trước khi lên lịch chơi.',
'Dự kiến có tính năng tùy chỉnh Memoria. Giữ Memoria Shards cho đến khi hệ thống cuối cùng được giải thích.',
'Bài tham khảo mô tả các lựa chọn sau trong phiên bản 14.5. Kiểm tra xem chúng còn mở trong game hay không.',
['Thử Lunar Bridge: 11 tầng, mỗi tầng có một trùm quỷ. Trùm có thể rơi Demon Familiar dành cho đơn vị Demon Controller.', 'Tra cứu đơn vị trong Codex của game. Gauntlets có vòng đấu loại 8 người; Odyssey cho chọn thẻ giữa các lượt chơi.', 'Giữ Trait Rerolls, Memoria Shards và vật liệu nâng cấp. Chờ cân bằng sát thương và hiệu ứng bị động Familiar cuối cùng trước khi dùng nhiều tài nguyên.'],
['Chạy game Android tương thích trên đám mây để giảm tải bộ nhớ, pin và phần cứng thiết bị chính khi chơi lâu.', 'Truy cập từ xa giúp quay lại phiên chơi và kiểm tra tiến độ. Bạn vẫn cần tự chơi; lưu trữ game trên đám mây không tự động farm.', 'Thử đăng nhập và chơi trước. Kết nối, nút máy chủ và khả năng tương thích của game ảnh hưởng đến trải nghiệm.'],
'Tích trữ vật liệu trước rồi chọn nâng cấp khi có chi tiết cuối cùng. Thử một phiên ngắn trên UGphone trước khi lên kế hoạch chơi lâu.',
'Cảnh trong đoạn giới thiệu Transcendent Sung Jin-woo', 'Nguồn ảnh', 'Hướng dẫn chuẩn bị', 'Giữ vật liệu và xem các tính năng kỷ niệm đã công bố trước khi chọn nâng cấp.']
}
TITLES={
'en':'Anime Vanguards 2nd Anniversary Preparation Guide',
'id':'Panduan Persiapan Ulang Tahun Kedua Anime Vanguards',
'pt':'Guia de Preparação para o Segundo Aniversário de Anime Vanguards',
'es':'Guía de preparación para el segundo aniversario de Anime Vanguards',
'th':'คู่มือเตรียมตัวครบรอบ 2 ปี Anime Vanguards',
'zh-TW':'Anime Vanguards 二週年改版準備攻略',
'vi':'Hướng dẫn chuẩn bị kỷ niệm 2 năm Anime Vanguards'}

for code,d in DATA.items():
    old=(ROOT/'outputs'/f'anime-vanguards-2nd-anniversary-guide-{code}.md').read_text(encoding='utf-8')
    chunks=re.split(r'\n## ',old)
    title=TITLES[code]
    sections=[x.split('\n',1) for x in chunks[1:]]
    meta=sections[0][1].strip()
    meta=re.sub(r'(?m)^- SEO Title:.*$',f'- SEO Title: {title}',meta)
    meta=re.sub(r'(?m)^- Meta Description:.*$',f'- Meta Description: {d[9]}',meta)
    if code=='id':
        meta=meta.replace('Anniversary ke-2','ulang tahun kedua').replace('anniversary','ulang tahun')
        meta=re.sub(r'(?m)^- Slug:.*$', '- Slug: anime-vanguards-ulang-tahun-kedua-guide-id', meta)
    meta=re.sub(r'(?m)^- Long-tail Keywords:.*$',lambda m:m.group().replace('anniversary','kỷ niệm') if code=='vi' else m.group(),meta)
    update=sections[1][1]
    bullets=[x[2:] for x in update.splitlines() if x.startswith('- ')]
    bullets.insert(0,d[1])
    # Keep the source's teaser claims, rather than treating them as a live reward.
    folder=OUT/code; (folder/'images').mkdir(parents=True,exist_ok=True)
    img=folder/'images'/'sung-jin-woo-preview.jpg'
    Image.open(r'C:\Users\Og\Downloads\anime-vanguards-2nd-anniversary-update-is-a-new-game-1025012901.webp').convert('RGB').save(img,quality=92)
    body=[('h',title),('p',d[0]),('h2',sections[1][0])]+[('b',x) for x in bullets]
    body += [('img',d[6]),('p',d[7]+': '+IMGURL),('h2',sections[2][0]),('p',d[2])]+[('n',x) for x in d[3]]
    body += [('h2',sections[3][0])]+[('b',x) for x in d[4]]
    # The play steps already form the preparation checklist; avoid duplicating them.
    body += [('h2',sections[5][0]),('p',d[5]),('h2',sections[6][0]),('p',SOURCE),('h2',sections[0][0])]
    body += [('p',x[2:]) for x in meta.splitlines()]
    md=[]; doc=Document(); sec=doc.sections[0]
    sec.top_margin=sec.bottom_margin=Inches(.65); sec.left_margin=sec.right_margin=Inches(.75)
    sec.page_height=Inches(11.7); sec.page_width=Inches(8.3)
    for sty in ['Normal','Title','Heading 1','Heading 2','Caption','List Bullet','List Number']:
        s=doc.styles[sty]; s.font.name='Arial'; s.font.color.rgb=RGBColor(0,0,0)
        s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei')
        s._element.rPr.rFonts.set(qn('w:cs'),'Tahoma')
    doc.styles['Normal'].font.size=Pt(10.5)
    doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    doc.styles['Normal'].paragraph_format.line_spacing=1.08
    doc.styles['Title'].font.size=Pt(22)
    doc.styles['Heading 1'].font.size=Pt(14)
    doc.styles['Heading 1'].paragraph_format.space_before=Pt(10)
    doc.styles['Caption'].font.size=Pt(9)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    doc.core_properties.title=title; doc.core_properties.language=code
    num=0
    for kind,text in body:
        if kind=='h': md.append('# '+text); doc.add_paragraph(text,'Title')
        elif kind=='h2': md.append('## '+text); doc.add_heading(text,1); num=0
        elif kind=='b': md.append('- '+text); doc.add_paragraph(text,'List Bullet')
        elif kind=='n': num+=1; md.append(f'{num}. '+text); doc.add_paragraph(text,'List Number')
        elif kind=='img':
            md.append(f'![{text}](images/sung-jin-woo-preview.jpg)\n\n*{text}*')
            p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(img),width=Inches(4.4))
            doc.add_paragraph(text,'Caption')
        else:
            md.append(text); p=doc.add_paragraph(text)
            if text.startswith(('http',d[7]+':')):
                for r in p.runs:r.font.size=Pt(8)
    stem='anime-vanguards-2nd-anniversary-guide-'+code
    (folder/(stem+'.md')).write_text('\n\n'.join(md)+'\n',encoding='utf-8')
    doc.save(folder/(stem+'.docx'))
    saved=Document(folder/(stem+'.docx'))
    assert saved.core_properties.title==title and saved.paragraphs[0].text==title
    assert len(saved.inline_shapes)==1
    if code!='zh-TW': assert not re.search('[\u4e00-\u9fff]', '\n'.join(p.text for p in saved.paragraphs))
print('Created and structurally checked 14 files')
