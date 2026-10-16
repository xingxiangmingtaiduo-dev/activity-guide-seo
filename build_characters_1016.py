import json, shutil, zipfile
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).parent
OUT=ROOT/'outputs'/'2026-10-16'/'anime-vanguards-new-characters-events'
ASSET=ROOT/'work'/'character-assets'
SOURCE='https://vanguards.gg/changelog'
slug='anime-vanguards-new-characters-events'
data=json.loads((ROOT/'work'/'character_articles.json').read_text(encoding='utf-8'))
codes=['en','id','pt','es','th','zh-TW','vi']

def hyperlink(p,text,url):
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); rp=OxmlElement('w:rPr'); c=OxmlElement('w:color'); c.set(qn('w:val'),'0563C1'); rp.append(c); r.append(rp)
    t=OxmlElement('w:t'); t.text=text; r.append(t); h.append(r); p._p.append(h)

def add_md(lines, s): lines.append(s)
def make(code):
    a=data[code]; folder=OUT/code; imgdir=folder/'images'; imgdir.mkdir(parents=True,exist_ok=True)
    shutil.copy2(ASSET/'new-units.jpg',imgdir/'new-units.jpg'); shutil.copy2(ASSET/'anniversary-banner.jpg',imgdir/'anniversary-banner.jpg')
    lines=['## '+a['heads'][4],'','- SEO Title: '+a['title'],'- Slug: '+slug+'-'+code.lower(),'- Primary Keyword: '+a['primary'],'- Long-tail Keywords: '+a['long'],'- Meta Description: '+a['meta'],'- Tags: Anime Vanguards, 15.0, Anniversary, '+a['primary'],'- Language: '+a['language'],'- Source: '+SOURCE,'- Version/Checked: '+a['date'],'','# '+a['title'],'',a['intro'],'','## '+a['heads'][0]]
    for x in a['characters']: lines += ['- '+x]
    lines += ['', '![%s](images/new-units.jpg)'%a['captions'][0], '*%s*'%a['captions'][0], '', '## '+a['heads'][1]]
    for x in a['events']: lines += ['- '+x]
    lines += ['', '![%s](images/anniversary-banner.jpg)'%a['captions'][1], '*%s*'%a['captions'][1], '', '## '+a['heads'][2]]
    for x in a['ug']: lines += ['- '+x]
    lines += ['', '## '+a['heads'][3]]
    for i,x in enumerate(a['checklist'],1): lines += [f'{i}. {x}']
    lines += ['', a['cta'], '', '## '+a['heads'][5],a['discord']+' ['+a['link']+'](https://discord.gg/'+('FhSaQfq6rJ' if code in ('pt','es') else 'Agkk96vcfA')+')']
    (folder/(slug+'-'+code.lower()+'.md')).write_text('\n\n'.join(lines)+'\n',encoding='utf-8')
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
    normal=doc.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(10.5); normal._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft JhengHei'); normal._element.rPr.rFonts.set(qn('w:cs'),'Tahoma'); normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.08
    title=doc.styles['Title']; title.font.name='Arial'; title.font.size=Pt(22); title.font.bold=True; title.font.color.rgb=RGBColor(0,0,0)
    for st in ('Heading 1','Heading 2'):
        doc.styles[st].font.name='Arial'; doc.styles[st].font.color.rgb=RGBColor(0,0,0)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    doc.core_properties.title=a['title']; doc.core_properties.language=code
    doc.add_paragraph(a['heads'][4],'Heading 1')
    for x in ['SEO Title: '+a['title'],'Slug: '+slug+'-'+code.lower(),'Primary Keyword: '+a['primary'],'Long-tail Keywords: '+a['long'],'Meta Description: '+a['meta'],'Tags: Anime Vanguards, 15.0, Anniversary, '+a['primary'],'Language: '+a['language'],'Source: '+SOURCE,'Version/Checked: '+a['date']]: doc.add_paragraph(x)
    doc.add_paragraph(a['title'],'Title'); doc.add_paragraph(a['intro'])
    def head(s):
        p=doc.add_paragraph(s,'Heading 1')
        if s in (a['heads'][1],a['heads'][2]): p.paragraph_format.page_break_before=True
    head(a['heads'][0]); [doc.add_paragraph(x,'List Bullet') for x in a['characters']]
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(imgdir/'new-units.jpg'),width=Inches(6.4)); q=doc.add_paragraph(a['captions'][0]); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    head(a['heads'][1]); [doc.add_paragraph(x,'List Bullet') for x in a['events']]
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(imgdir/'anniversary-banner.jpg'),width=Inches(5.4)); q=doc.add_paragraph(a['captions'][1]); q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    head(a['heads'][2]); [doc.add_paragraph(x,'List Bullet') for x in a['ug']]
    head(a['heads'][3]); [doc.add_paragraph(x,'List Number') for x in a['checklist']]; doc.add_paragraph(a['cta'])
    head(a['heads'][5]); doc.add_paragraph(a['discord']); p=doc.add_paragraph(); hyperlink(p,a['link'],'https://discord.gg/'+('FhSaQfq6rJ' if code in ('pt','es') else 'Agkk96vcfA'))
    doc.save(folder/(slug+'-'+code.lower()+'.docx'))

for c in codes: make(c)
zip_path=OUT/(slug+'.zip')
with zipfile.ZipFile(zip_path,'w',zipfile.ZIP_DEFLATED) as z:
    for p in OUT.rglob('*'):
        if p.is_file() and p != zip_path: z.write(p,p.relative_to(OUT))
print('created',len(codes)*2,zip_path)

